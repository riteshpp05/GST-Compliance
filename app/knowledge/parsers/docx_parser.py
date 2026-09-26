"""
app.knowledge.parsers.docx_parser
=================================
DOCX Document Parser implementation for UC15 Knowledge Subsystem (Sprint 12.3).
Extracts paragraphs and heading sections using python-docx.
"""

from __future__ import annotations

import io
from typing import Optional, Union
from app.knowledge.exceptions import DocumentParsingError
from app.knowledge.models import KnowledgeDocument, KnowledgeMetadata
from app.knowledge.parsers.base import DocumentParser

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class DOCXParser(DocumentParser):
    """Parses DOCX documents into KnowledgeDocument objects."""

    def parse(
        self,
        file_source: Union[str, bytes],
        filename: str,
        metadata: Optional[KnowledgeMetadata] = None,
    ) -> KnowledgeDocument:
        import os

        file_bytes = file_source if isinstance(file_source, bytes) else None
        target_name = file_source if (isinstance(file_source, str) and os.path.exists(file_source)) else filename
        self.validate_file(target_name, file_bytes)

        if not DOCX_AVAILABLE:
            raise DocumentParsingError("python-docx library is not installed. DOCX parsing unavailable.")

        meta = self._default_metadata(filename, metadata)
        paragraphs_text = []

        try:
            if isinstance(file_source, bytes):
                stream = io.BytesIO(file_source)
                doc = docx.Document(stream)
            else:
                doc = docx.Document(file_source)

            for p in doc.paragraphs:
                txt = p.text.strip()
                if not txt:
                    continue
                if p.style and p.style.name and p.style.name.startswith("Heading"):
                    paragraphs_text.append(f"\n### Section: {txt}\n")
                else:
                    paragraphs_text.append(txt)

        except Exception as exc:
            raise DocumentParsingError(f"Failed to parse DOCX document '{filename}': {str(exc)}") from exc

        if not paragraphs_text:
            raise DocumentParsingError(f"DOCX document '{filename}' contains no extractable text.")

        full_content = "\n".join(paragraphs_text)
        return KnowledgeDocument(
            id=meta.document_id,
            title=meta.document_name,
            metadata=meta,
            raw_content=full_content,
            file_path=file_source if isinstance(file_source, str) else None,
        )
