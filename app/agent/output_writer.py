"""
UC15 GST Compliance Agent — Results Writer
Generates Excel workbooks with formatted audit sheets and structured JSON artifacts.
"""
from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config.settings import settings
from app.domain.models.validation import ComplianceDecision
from app.infrastructure.logging import get_logger

logger = get_logger(__name__)

NAVY = "1C2B3A"
WHITE = "FFFFFF"
STATUS_FILL = {
    "COMPLIANT": "D9F2E3",
    "NEEDS_REVIEW": "FDF3D0",
    "NON_COMPLIANT": "FBE0E0",
    "BLOCKED": "FBE0E0",
}


class ResultsWriter:
    """Writes compliance decisions into styled Excel reports and JSON run artifacts."""

    def __init__(self, results_dir: Optional[str] = None):
        self.results_dir = results_dir or settings.results_dir
        os.makedirs(self.results_dir, exist_ok=True)

    def write_results(self, decisions: List[ComplianceDecision]) -> Dict[str, Any]:
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        xlsx_path = os.path.join(self.results_dir, f"uc15_gst_run_{ts}.xlsx")
        json_path = os.path.join(self.results_dir, f"uc15_gst_run_{ts}.json")

        self._write_excel(decisions, xlsx_path)
        self._write_json(decisions, json_path, ts)
        logger.info(f"Results written to {xlsx_path} and {json_path}")
        return {
            "excel": xlsx_path,
            "json": json_path,
            "run_id": ts,
            "invoice_count": len(decisions),
        }

    def _write_excel(self, decisions: List[ComplianceDecision], path: str) -> None:
        writer = pd.ExcelWriter(path, engine="openpyxl")
        try:
            self._sheet_summary(decisions, writer)
            self._sheet_detailed(decisions, writer)
        finally:
            writer.close()
        self._style_workbook(path)

    def _sheet_summary(self, decisions: List[ComplianceDecision], writer: pd.ExcelWriter) -> None:
        total = len(decisions)
        by_status = pd.Series([d.status for d in decisions]).value_counts()
        rows = [
            ("Total invoices validated", total),
            ("Compliant — filing ready", int(by_status.get("COMPLIANT", 0))),
            ("Needs review", int(by_status.get("NEEDS_REVIEW", 0))),
            ("Non-compliant — blocked", int(by_status.get("NON_COMPLIANT", 0))),
            ("", ""),
            (
                "GSTR filing readiness rate",
                f"{100 * by_status.get('COMPLIANT', 0) / total:.0f}%" if total else "0%",
            ),
            (
                "Average risk score",
                f"{sum(d.risk_score for d in decisions if d.risk_score is not None) / total:.1f}" if total else "0.0",
            ),
            (
                "High / Critical risk count",
                sum(1 for d in decisions if getattr(d, "risk_level", None) in ("HIGH", "CRITICAL")),
            ),
            ("Run generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")),
        ]
        pd.DataFrame(rows, columns=["Metric", "Value"]).to_excel(
            writer, sheet_name="Summary", index=False
        )

    def _sheet_detailed(self, decisions: List[ComplianceDecision], writer: pd.ExcelWriter) -> None:
        df = pd.DataFrame([
            {
                "Invoice": d.invoice_no,
                "Direction": d.direction,
                "Counterparty": d.counterparty_name,
                "GSTIN": d.counterparty_gstin,
                "HSN": d.hsn_code,
                "Taxable Value": d.taxable_value_inr,
                "Gates Failed": d.failed_gate_count,
                "Status": d.status,
                "Risk Score": d.risk_score if d.risk_score is not None else 0.0,
                "Risk Level": getattr(d, "risk_level", "LOW"),
                "Priority": getattr(d, "priority", "P4"),
                "Justification": d.justification,
                "SAP Action": d.sap_action,
                "Audit Ref": d.audit_trail_ref,
            }
            for d in decisions
        ])
        df.to_excel(writer, sheet_name="Detailed Results", index=False)


    def _style_workbook(self, path: str) -> None:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path)
        except Exception as e:
            logger.warning(f"Could not open workbook {path} for styling: {e}")
            return

        header_font = Font(bold=True, color=WHITE)
        header_fill = PatternFill("solid", fgColor=NAVY)

        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for col_cells in ws.columns:
                length = max(
                    (len(str(c.value)) if c.value is not None else 0) for c in col_cells
                )
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(
                    45, max(12, length + 2)
                )
            ws.freeze_panes = "A2"

            if ws.title == "Detailed Results":
                status_col = None
                for cell in ws[1]:
                    if cell.value == "Status":
                        status_col = cell.column
                        break
                if status_col:
                    for row in ws.iter_rows(min_row=2, min_col=status_col, max_col=status_col):
                        for cell in row:
                            fill = STATUS_FILL.get(cell.value)
                            if fill:
                                for c in ws[cell.row]:
                                    c.fill = PatternFill("solid", fgColor=fill)
        wb.save(path)

    def _write_json(self, decisions: List[ComplianceDecision], path: str, run_id: str) -> None:
        results_data = []
        for d in decisions:
            if hasattr(d, "to_dict"):
                results_data.append(d.to_dict())
            elif hasattr(d, "__dict__"):
                from dataclasses import asdict
                try:
                    results_data.append(asdict(d))
                except TypeError:
                    results_data.append(d.__dict__)
            else:
                results_data.append(dict(d))

        payload = {
            "run_id": run_id,
            "generated_at_utc": datetime.now(timezone.utc).isoformat(),
            "invoice_count": len(decisions),
            "results": results_data,
        }
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2, default=str)
