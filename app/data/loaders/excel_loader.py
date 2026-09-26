"""
UC15 GST Compliance Agent — Excel Data Loader (v2.0)
Loads Section A (Invoices), Section B (HSN Master), Section C (State Codes) from Excel dataset.
Implements non-crashing batch ingestion with rich IngestionBatchResult reporting.
"""
from __future__ import annotations

import os
from decimal import Decimal
from typing import Dict, List, Optional
import openpyxl

from app.config.settings import (
    EXCEL_FILE,
    SHEET_NAME,
    SEC_A_START,
    SEC_A_END,
    SEC_B_START,
    SEC_B_END,
    SEC_C_START,
    SEC_C_END,
)
from app.data.loaders.base import BaseInvoiceLoader
from app.data.normalization.invoice_normalizer import InvoiceNormalizer
from app.domain.exceptions import DataLoadError
from app.domain.models.ingestion import (
    IngestionBatchResult,
    IngestionRecordResult,
    IngestionStatus,
    SourceMetadata,
)
from app.domain.models.invoice import Invoice
from app.domain.models.tax import HSNMaster
from app.infrastructure.logging import get_logger

logger = get_logger(__name__)


class ExcelInvoiceLoader(BaseInvoiceLoader):
    """Loads invoices and reference data from the multi-section Excel workbook."""

    def __init__(self, excel_path: Optional[str] = None, sheet_name: Optional[str] = None):
        self.excel_path = excel_path or EXCEL_FILE
        self.sheet_name = sheet_name or SHEET_NAME
        self._batch_result: Optional[IngestionBatchResult] = None
        self._hsn_master: Dict[str, HSNMaster] = {}
        self._state_codes: Dict[str, str] = {}
        self._loaded = False

    def _ensure_loaded(self) -> None:
        if self._loaded and self._batch_result is not None:
            return

        if not os.path.exists(self.excel_path):
            raise FileNotFoundError(f"Excel dataset not found: {self.excel_path}")

        logger.info(f"Loading workbook: {self.excel_path} (sheet: {self.sheet_name})")
        try:
            wb = openpyxl.load_workbook(self.excel_path, data_only=True)
            ws = wb[self.sheet_name] if self.sheet_name in wb.sheetnames else wb.active
            end_row = max(SEC_A_END, ws.max_row)

            # Detect Header Row dynamically
            header_map: Dict[str, int] = {}
            header_row_idx = None
            all_rows = list(ws.iter_rows(min_row=1, max_row=min(10, ws.max_row), values_only=True))

            for idx, r in enumerate(all_rows):
                if not r:
                    continue
                row_strs = [str(cell).strip().lower() for cell in r if cell is not None]
                if any("invoice" in s for s in row_strs) and any("gstin" in s or "counterparty" in s for s in row_strs):
                    header_row_idx = idx + 1  # 1-based index
                    for col_idx, cell in enumerate(r):
                        if cell is None:
                            continue
                        clean_col = "".join(ch for ch in str(cell).lower() if ch.isalnum())
                        header_map[clean_col] = col_idx
                    break

            start_data_row = (header_row_idx + 1) if header_row_idx else SEC_A_START

            # Helper for header lookup
            def get_col_val(r: tuple, keys: List[str], default_idx: int) -> Any:
                for k in keys:
                    if k in header_map and header_map[k] < len(r):
                        return r[header_map[k]]
                if default_idx < len(r):
                    return r[default_idx]
                return None

            # 1. Invoices (Section A)
            records: List[IngestionRecordResult] = []
            row_idx = 0
            for row in ws.iter_rows(min_row=start_data_row, max_row=end_row, values_only=True):
                row_idx += 1
                if not row or row[0] is None or str(row[0]).strip() == "":
                    continue

                inv_id = str(row[0]).strip()
                inv_id_lower = inv_id.lower()
                if inv_id_lower.startswith("invoice no") or inv_id_lower == "invoice_no" or inv_id_lower.startswith("section"):
                    continue  # skip header row if hit
                if inv_id_lower.startswith("total") or inv_id_lower.startswith("hsn") or inv_id_lower.startswith("state") or inv_id_lower.startswith("test"):
                    continue  # skip non-invoice section headers or reference table rows

                meta = SourceMetadata(
                    source_type="excel",
                    source_name="ExcelSectionA",
                    source_file=self.excel_path,
                    source_sheet=self.sheet_name,
                    source_record_id=inv_id,
                )

                raw_dict = {
                    "invoice_no": get_col_val(row, ["invoiceno", "invoicenumber", "invoiceid"], 0),
                    "invoice_date": get_col_val(row, ["invoicedate", "date"], 1),
                    "direction": get_col_val(row, ["direction"], 2),
                    "counterparty_gstin": get_col_val(row, ["counterpartygstin", "gstin"], 3),
                    "counterparty_name": get_col_val(row, ["counterpartyname", "vendorname", "customername", "party"], 4),
                    "place_of_supply": get_col_val(row, ["placeofsupply", "pos"], 5),
                    "hsn_code": get_col_val(row, ["hsncode", "hsnsac", "hsn"], 6),
                    "item_desc": get_col_val(row, ["itemdesc", "description"], 7),
                    "taxable_value_inr": get_col_val(row, ["taxablevalueinr", "taxablevalue"], 8),
                    "cgst_rate": get_col_val(row, ["cgst", "cgstrate"], 9),
                    "sgst_rate": get_col_val(row, ["sgst", "sgstrate"], 10),
                    "igst_rate": get_col_val(row, ["igst", "igstrate"], 11),
                    "total_tax_inr": get_col_val(row, ["totaltaxinr", "totaltax"], 12),
                    "total_amt": get_col_val(row, ["totalamt", "totalamount"], 13 if header_map else (13 if len(row) > 15 else 12)),
                    "eway_bill_status": get_col_val(row, ["ewaybillstatus", "ewaybill"], 14 if header_map else (14 if len(row) > 15 else 13)),
                    "gstr2b_reflected": get_col_val(row, ["gstr2breflected", "gstr2b"], 15 if header_map else (15 if len(row) > 15 else 14)),
                }

                rec_result = InvoiceNormalizer.normalize_record(
                    raw_dict,
                    source_metadata=meta,
                    record_index=row_idx,
                )
                records.append(rec_result)

            self._batch_result = IngestionBatchResult(
                source_name=f"Excel:{os.path.basename(self.excel_path)}",
                total_records=len(records),
                records=records,
            )

            # 2. HSN Master (Section B or dedicated tab)
            self._hsn_master = {}
            ws_hsn = wb["HSN_Master"] if "HSN_Master" in wb.sheetnames else ws
            hsn_min_row = 2 if "HSN_Master" in wb.sheetnames else SEC_B_START
            hsn_max_row = ws_hsn.max_row if "HSN_Master" in wb.sheetnames else SEC_B_END
            for row in ws_hsn.iter_rows(min_row=hsn_min_row, max_row=hsn_max_row, values_only=True):
                if not row or row[0] is None or str(row[0]).strip() == "":
                    continue
                try:
                    hsn = HSNMaster.from_raw(
                        hsn_code=str(row[0]),
                        description=str(row[1]) if len(row) > 1 else "",
                        correct_cgst_rate=row[2] if len(row) > 2 else 0,
                        correct_sgst_rate=row[3] if len(row) > 3 else 0,
                        correct_igst_rate=row[4] if len(row) > 4 else 0,
                    )
                    self._hsn_master[hsn.hsn_code] = hsn
                except Exception:
                    continue

            # 3. State Codes (Section C or dedicated tab)
            self._state_codes = {}
            ws_st = wb["State_Codes"] if "State_Codes" in wb.sheetnames else ws
            st_min_row = 2 if "State_Codes" in wb.sheetnames else SEC_C_START
            st_max_row = ws_st.max_row if "State_Codes" in wb.sheetnames else SEC_C_END
            for row in ws_st.iter_rows(min_row=st_min_row, max_row=st_max_row, values_only=True):
                if not row or row[0] is None or str(row[0]).strip() == "":
                    continue
                self._state_codes[str(row[0]).strip()] = str(row[1]).strip() if len(row) > 1 else str(row[0]).strip()

            self._loaded = True
            logger.info(
                f"Successfully ingested {self._batch_result.valid_count}/{self._batch_result.total_records} "
                f"valid invoices, {len(self._hsn_master)} HSN masters, and {len(self._state_codes)} state codes."
            )
        except FileNotFoundError:
            raise
        except Exception as e:
            raise DataLoadError(f"Failed loading workbook {self.excel_path}: {e}") from e

    def load(self) -> IngestionBatchResult:
        self._ensure_loaded()
        assert self._batch_result is not None
        return self._batch_result

    def load_invoices(self) -> List[Invoice]:
        return self.load().valid_invoices

    def load_hsn_master(self) -> Dict[str, HSNMaster]:
        self._ensure_loaded()
        return dict(self._hsn_master)

    def load_state_codes(self) -> Dict[str, str]:
        self._ensure_loaded()
        return dict(self._state_codes)
