"""
app.data.sources.excel_source
=============================
Excel Data Source Adapter for UC15 (Sprint 17).
Supports reading from file path or openpyxl workbook buffer.
"""

from __future__ import annotations

import io
from typing import Any, Dict, List, Optional, Union
import openpyxl

from app.data.sources.base import DataSource, DataSourceMetadata, DataSourceType


class ExcelDataSource(DataSource):
    """Concrete Excel Data Source Adapter."""

    def __init__(
        self,
        source_id: str,
        source_name: str,
        file_path: Optional[str] = None,
        excel_bytes: Optional[bytes] = None,
        sheet_name: Optional[str] = None,
    ):
        self.source_id = source_id
        self.source_name = source_name
        self.file_path = file_path
        self.excel_bytes = excel_bytes
        self.sheet_name = sheet_name
        self._connected = False

    def connect(self) -> bool:
        if self.file_path:
            import os
            self._connected = os.path.exists(self.file_path)
            return self._connected
        elif self.excel_bytes is not None:
            self._connected = True
            return True
        return False

    def fetch_records(self) -> List[Dict[str, Any]]:
        if not self._connected and not self.connect():
            raise ValueError(f"ExcelDataSource '{self.source_name}' is not connected or accessible.")

        if self.file_path:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
        elif self.excel_bytes:
            wb = openpyxl.load_workbook(io.BytesIO(self.excel_bytes), data_only=True)
        else:
            return []

        sheet = wb[self.sheet_name] if self.sheet_name and self.sheet_name in wb.sheetnames else wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h).strip() if h is not None else f"col_{idx}" for idx, h in enumerate(rows[0])]
        records: List[Dict[str, Any]] = []

        for r_idx, row in enumerate(rows[1:], start=2):
            if not any(row):
                continue
            record_dict = {}
            for c_idx, val in enumerate(row):
                if c_idx < len(headers):
                    key = headers[c_idx]
                    record_dict[key] = val
            records.append(record_dict)

        return records

    def metadata(self) -> DataSourceMetadata:
        return DataSourceMetadata(
            source_id=self.source_id,
            source_type=DataSourceType.EXCEL,
            source_name=self.source_name,
            source_location=self.file_path or "IN_MEMORY_EXCEL",
            additional_info={"sheet_name": self.sheet_name},
        )
