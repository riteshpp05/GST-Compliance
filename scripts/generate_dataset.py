#!/usr/bin/env python3
"""
Generates data/UC15_GSTCompliance_Dataset.xlsx — a 4-section demo dataset
for the UC15 GST & Tax Compliance Validation Agent.

Row placement is driven directly by config/settings.py's SEC_*_HEADER /
SEC_*_START / SEC_*_END constants. Invoices are deliberately constructed
per gate-failure count (0, 1, or 2+ failures) rather than randomized,
since this UC validates a checklist, not a composite score — the whole
point is being able to demonstrate each gate failing independently.
"""
import os
import sys
import random
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config.settings import (
    EXCEL_FILE, SHEET_NAME,
    SEC_A_HEADER, SEC_A_START, SEC_A_END,
    SEC_B_HEADER, SEC_B_START, SEC_B_END,
    SEC_C_HEADER, SEC_C_START, SEC_C_END,
    SEC_D_HEADER, SEC_D_START, SEC_D_END,
)

random.seed(15)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(HERE, "..", EXCEL_FILE)

NAVY = "1C2B3A"
LIGHT = "EEF2F7"
WHITE = "FFFFFF"

HSN_MASTER = [
    ("8409", "Parts for Internal Combustion Engines", 9, 9, 18),
    ("8483", "Transmission Shafts & Cranks", 9, 9, 18),
    ("8481", "Taps, Cocks, Valves", 9, 9, 18),
    ("3926", "Plastic Articles NES", 9, 9, 18),
    ("7318", "Screws, Bolts, Nuts (Iron/Steel)", 9, 9, 18),
    ("8536", "Electrical Switches & Connectors", 9, 9, 18),
    ("4819", "Cartons, Boxes, Cases (Paper)", 6, 6, 12),
    ("8501", "Electric Motors & Generators", 9, 9, 18),
    ("8544", "Insulated Wire & Cable", 9, 9, 18),
    ("9983", "Other Professional/Technical Services", 9, 9, 18),
    ("9954", "Construction Services", 9, 9, 18),
    ("8708", "Motor Vehicle Parts & Accessories", 14, 14, 28),
    ("3923", "Plastic Packing Articles", 9, 9, 18),
    ("8607", "Railway/Tramway Parts", 9, 9, 18),
    ("8414", "Pumps, Compressors, Fans", 9, 9, 18),
]

STATE_CODES = [
    ("27", "Maharashtra"), ("07", "Delhi"), ("29", "Karnataka"), ("33", "Tamil Nadu"),
    ("06", "Haryana"), ("24", "Gujarat"), ("09", "Uttar Pradesh"), ("19", "West Bengal"),
    ("36", "Telangana"), ("32", "Kerala"),
]

COUNTERPARTIES = [
    "Bosch India Ltd", "Bajaj Auto Ltd", "Hero MotoCorp", "Maruti Suzuki",
    "Tata Motors", "Ashok Leyland", "Motherson Sumi", "Minda Industries",
    "Sundram Fasteners", "Endurance Technologies", "Varroc Engineering",
    "Bharat Forge", "Exide Industries", "Amara Raja Batteries", "JBM Auto",
]


def make_gstin(state_code, seed_suffix, malformed=False):
    if malformed:
        return f"{state_code}INVALID{seed_suffix}"
    pan_letters = "".join(random.choices("ABCDEFGHIJKLMNOPQRSTUVWXYZ", k=5))
    pan_digits = f"{random.randint(1000, 9999)}"
    pan_check = random.choice("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    entity_code = str(random.randint(1, 9))
    checksum = random.choice("0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    return f"{state_code}{pan_letters}{pan_digits}{pan_check}{entity_code}Z{checksum}"


def rdate(base_days_offset, spread=0):
    d = datetime.now() + timedelta(days=base_days_offset + random.randint(0, spread))
    return d.strftime("%Y-%m-%d")


def section_title(ws, row, span, title):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = ws.cell(row=row, column=1, value=title)
    c.font = Font(name="Calibri", bold=True, size=13, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24


def col_hdr(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
        c.fill = PatternFill("solid", fgColor="34486B")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 28


def data_row(ws, row, values, shaded):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name="Calibri", size=10)
        if shaded:
            c.fill = PatternFill("solid", fgColor=LIGHT)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    section_title(ws, 1, 15, "USE CASE 15 - GST & TAX COMPLIANCE VALIDATION AGENT")

    section_title(ws, SEC_B_HEADER - 1, 5, "SECTION B: HSN/SAC MASTER WITH CORRECT RATES")
    col_hdr(ws, SEC_B_HEADER, ["HSN Code", "Description", "Correct CGST %", "Correct SGST %", "Correct IGST %"],
            [10, 34, 14, 14, 14])
    n_b = SEC_B_END - SEC_B_START + 1
    for i in range(min(n_b, len(HSN_MASTER))):
        row_no = SEC_B_START + i
        data_row(ws, row_no, list(HSN_MASTER[i]), row_no % 2 == 0)

    section_title(ws, SEC_C_HEADER - 1, 2, "SECTION C: GSTIN STATE CODE REFERENCE")
    col_hdr(ws, SEC_C_HEADER, ["State Code", "State Name"], [12, 20])
    n_c = SEC_C_END - SEC_C_START + 1
    for i in range(min(n_c, len(STATE_CODES))):
        row_no = SEC_C_START + i
        data_row(ws, row_no, list(STATE_CODES[i]), row_no % 2 == 0)

    col_hdr(ws, SEC_A_HEADER, ["Invoice No", "Invoice Date", "Direction", "Counterparty GSTIN",
                               "Counterparty Name", "Place of Supply", "HSN Code", "Item Desc",
                               "Taxable Value (INR)", "CGST %", "SGST %", "IGST %", "Total Amt",
                               "E-Way Bill Status", "GSTR-2B Reflected"],
            [14, 12, 10, 18, 22, 16, 10, 30, 16, 8, 8, 8, 16, 16, 16])

    n_a = SEC_A_END - SEC_A_START + 1
    plan = (["CLEAN"] * 15 + ["ONE_FAIL"] * 9 + ["GSTIN_FAIL"] * 2 + ["MULTI_FAIL"] * 4)
    random.shuffle(plan)

    for i in range(n_a):
        row_no = SEC_A_START + i
        kind = plan[i]
        hsn_code, hsn_desc, cgst_m, sgst_m, igst_m = random.choice(HSN_MASTER)
        state_code, state_name = random.choice(STATE_CODES)
        direction = random.choice(["AR", "AP"])
        counterparty = random.choice(COUNTERPARTIES)
        taxable_value = random.choice([12000, 35000, 62000, 88000, 145000, 210000])
        same_state_invoice = random.random() < 0.5

        gstin = make_gstin(state_code, i, malformed=False)
        place_of_supply = state_name if same_state_invoice else random.choice(
            [s[1] for s in STATE_CODES if s[1] != state_name])

        if same_state_invoice:
            cgst, sgst, igst = cgst_m, sgst_m, 0
        else:
            cgst, sgst, igst = 0, 0, igst_m
        eway_status = ""
        gstr2b = True
        item_desc = f"{hsn_desc} - supply per contract"

        if kind == "CLEAN":
            if taxable_value > 50000:
                eway_status = "GENERATED"

        elif kind == "ONE_FAIL":
            # "hsn" deliberately excluded from this list — an unknown HSN
            # code cascades into both Gate 2 AND Gate 3 failing (you can't
            # verify a tax rate for a code you can't find), which is
            # correct real-world behavior, not a single-gate scenario.
            fail_choice = random.choice(["rate", "pos", "eway", "itc"])
            if fail_choice == "rate":
                if same_state_invoice:
                    cgst, sgst = cgst + 5, sgst + 5
                else:
                    igst = igst + 5
                direction = "AR"
                taxable_value = 35000
            elif fail_choice == "pos":
                if same_state_invoice:
                    cgst, sgst, igst = 0, 0, igst_m
                else:
                    cgst, sgst, igst = cgst_m, sgst_m, 0
                direction = "AR"
                taxable_value = 35000
            elif fail_choice == "eway":
                taxable_value = 88000
                eway_status = random.choice(["PENDING", "CANCELLED", ""])
                direction = "AR"
            elif fail_choice == "itc":
                direction = "AP"
                taxable_value = 35000
                gstr2b = False

        elif kind == "GSTIN_FAIL":
            gstin = make_gstin(state_code, i, malformed=True)
            taxable_value = 35000

        elif kind == "MULTI_FAIL":
            hsn_code = "9999"
            direction = "AP"
            taxable_value = 88000
            eway_status = "PENDING"
            gstr2b = False
            if same_state_invoice:
                cgst, sgst, igst = 0, 0, 18

        total_amt = round(taxable_value * (1 + (cgst + sgst + igst) / 100), 2)

        vals = [
            f"INV-{8000000+i}", rdate(-random.randint(1, 20)), direction, gstin, counterparty,
            place_of_supply, hsn_code, item_desc, taxable_value, cgst, sgst, igst, total_amt,
            eway_status, "YES" if gstr2b else "NO",
        ]
        data_row(ws, row_no, vals, row_no % 2 == 0)

    section_title(ws, SEC_D_HEADER - 1, 4,
                  "SECTION D: AGENT OUTPUT REFERENCE (prior-period actuals, for Learn-phase comparison - not agent input)")
    col_hdr(ws, SEC_D_HEADER, ["Invoice No (Historical)", "Final Status", "Gates Failed", "Resolution Notes"],
            [16, 16, 12, 30])
    n_d = SEC_D_END - SEC_D_START + 1
    for i in range(n_d):
        row_no = SEC_D_START + i
        status = random.choices(["COMPLIANT", "NEEDS_REVIEW", "NON_COMPLIANT"], weights=[0.5, 0.3, 0.2])[0]
        gates_failed = {"COMPLIANT": 0, "NEEDS_REVIEW": 1, "NON_COMPLIANT": random.randint(2, 3)}[status]
        note = {"COMPLIANT": "Filed in GSTR-1 without correction", "NEEDS_REVIEW": "Corrected before filing deadline",
                "NON_COMPLIANT": "Held, multi-issue correction required"}[status]
        data_row(ws, row_no, [f"INV-{7900000+i}", status, gates_failed, note], row_no % 2 == 0)

    ws.freeze_panes = f"A{SEC_A_HEADER + 1}"
    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)

    print(f"Wrote {OUT_PATH}")
    print(f"  Section A (Invoices):          {n_a} rows (rows {SEC_A_START}-{SEC_A_END})")
    print(f"  Section B (HSN Master):        {min(n_b, len(HSN_MASTER))} rows (rows {SEC_B_START}-{SEC_B_END})")
    print(f"  Section C (State Codes):       {min(n_c, len(STATE_CODES))} rows (rows {SEC_C_START}-{SEC_C_END})")
    print(f"  Section D (Output Reference):  {n_d} rows (rows {SEC_D_START}-{SEC_D_END})")
    print(f"  Designed mix: 15 clean, 9 one-fail, 2 GSTIN-fail, 4 multi-fail")


if __name__ == "__main__":
    main()
